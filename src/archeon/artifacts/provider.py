"""Lazy, verified artifact creation and editing providers."""

from __future__ import annotations

import csv
import ast
import hashlib
import html
import io
import json
import re
import shutil
import struct
import textwrap
import time
import zipfile
from pathlib import Path
from typing import Any

from .models import ArtifactResult, ArtifactSpec


class ArtifactProvider:
    def create(self, spec: ArtifactSpec) -> ArtifactResult:
        if spec.path.exists():
            raise FileExistsError("artifact_target_exists")
        if not spec.path.parent.is_dir():
            raise FileNotFoundError("artifact_parent_missing")
        getattr(self, f"_create_{spec.format}")(spec)
        result = self.verify(spec.path)
        expected_pictures = int(spec.metadata.get("expected_picture_count", 0) or 0)
        if spec.format == "pptx" and expected_pictures > int(result.structure.get("pictures", 0)):
            spec.path.unlink(missing_ok=True)
            raise ValueError("presentation_expected_images_not_embedded")
        return result

    def inspect(self, path: str | Path) -> dict[str, Any]:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise FileNotFoundError("artifact_not_found")
        artifact_format = source.suffix.lstrip(".").casefold()
        method = getattr(self, f"_inspect_{artifact_format}", None)
        if method is None:
            raise ValueError("unsupported_artifact_format")
        return method(source)

    def verify(self, path: str | Path) -> ArtifactResult:
        source = Path(path).expanduser().resolve()
        structure = self.inspect(source)
        size = source.stat().st_size
        verified = size > 0 and bool(structure.get("valid"))
        return ArtifactResult(
            str(source), source.suffix.lstrip(".").casefold(), size,
            hashlib.sha256(source.read_bytes()).hexdigest(), verified, structure,
        )

    def edit_copy(
        self, source: str | Path, output: str | Path, *, replacements: dict[str, str] | None = None,
        append: str = "", cell_updates: dict[str, Any] | None = None,
    ) -> ArtifactResult:
        source_path = Path(source).expanduser().resolve()
        output_path = Path(output).expanduser().resolve()
        if not source_path.is_file() or output_path.exists() or not output_path.parent.is_dir():
            raise ValueError("artifact_edit_paths_invalid")
        if source_path.suffix.casefold() != output_path.suffix.casefold():
            raise ValueError("artifact_edit_format_mismatch")
        checkpoint = source_path.with_name(f".{source_path.name}.checkpoint-{time.time_ns()}")
        shutil.copy2(source_path, checkpoint)
        artifact_format = source_path.suffix.lstrip(".").casefold()
        try:
            getattr(self, f"_edit_{artifact_format}")(
                source_path, output_path, replacements or {}, append, cell_updates or {},
            )
            verified = self.verify(output_path)
            return ArtifactResult(
                verified.path, verified.format, verified.bytes, verified.sha256,
                verified.verified, verified.structure, str(checkpoint),
            )
        except BaseException:
            output_path.unlink(missing_ok=True)
            raise

    @staticmethod
    def _create_txt(spec: ArtifactSpec) -> None:
        spec.path.write_text(spec.content, encoding="utf-8")

    _create_md = _create_txt
    _create_css = _create_txt
    _create_js = _create_txt
    _create_py = _create_txt

    @staticmethod
    def _create_json(spec: ArtifactSpec) -> None:
        value = spec.metadata.get("data")
        if value is None:
            value = json.loads(spec.content or "{}")
        spec.path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def _create_csv(spec: ArtifactSpec) -> None:
        with spec.path.open("x", encoding="utf-8-sig", newline="") as output:
            csv.writer(output).writerows(spec.rows)

    @staticmethod
    def _create_html(spec: ArtifactSpec) -> None:
        if re.match(r"\s*<!doctype\s+html\b", spec.content, flags=re.I):
            spec.path.write_text(spec.content, encoding="utf-8")
            return
        sections = "".join(
            f"<section><h2>{html.escape(str(item.get('heading', '')))}</h2><p>{html.escape(str(item.get('body', '')))}</p></section>"
            for item in spec.sections
        )
        document = (
            "<!doctype html><html lang=\"es\"><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width\">"
            f"<title>{html.escape(spec.title)}</title><body><main><h1>{html.escape(spec.title)}</h1>"
            f"<p>{html.escape(spec.content)}</p>{sections}</main></body></html>"
        )
        spec.path.write_text(document, encoding="utf-8")

    @staticmethod
    def _create_docx(spec: ArtifactSpec) -> None:
        from docx import Document
        from archeon.documents.style import DocumentStyleProfile

        document = Document()
        DocumentStyleProfile.from_mapping(spec.metadata.get("style_profile")).apply_docx(document)
        if spec.title:
            document.add_paragraph(spec.title, style="Title")
        if spec.content:
            document.add_paragraph(spec.content)
        for section in spec.sections:
            if section.get("page_break"):
                document.add_page_break()
            if section.get("heading"):
                document.add_heading(str(section["heading"]), level=max(1, min(int(section.get("level", 1)), 9)))
            if section.get("body"):
                paragraph = document.add_paragraph(str(section["body"]))
                if str(section.get("alignment", "")).casefold() == "left":
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            rows = section.get("rows") or []
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        table.cell(row_index, column_index).text = str(value)
            image = section.get("image")
            if image:
                from docx.shared import Inches
                paragraph = document.add_paragraph(); paragraph.alignment = 1
                paragraph.add_run().add_picture(str(image), width=Inches(float(section.get("image_width", 6.0))))
                if section.get("caption"):
                    caption = document.add_paragraph(str(section["caption"])); caption.alignment = 1
        document.core_properties.title = spec.title
        document.save(spec.path)

    @staticmethod
    def _create_xlsx(spec: ArtifactSpec) -> None:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = spec.sheets or [{"name": spec.title or "Datos", "rows": spec.rows}]
        for sheet_spec in sheets:
            sheet = workbook.create_sheet(str(sheet_spec.get("name", "Datos"))[:31])
            for row in sheet_spec.get("rows", []):
                sheet.append(list(row))
            if sheet.max_row:
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="007C91")
            for cell, formula in dict(sheet_spec.get("formulas", {})).items():
                sheet[str(cell)] = str(formula)
            chart_spec = sheet_spec.get("chart")
            if chart_spec and sheet.max_row >= 2 and sheet.max_column >= 2:
                chart = BarChart()
                data = Reference(sheet, min_col=int(chart_spec.get("data_col", 2)), min_row=1, max_row=sheet.max_row)
                categories = Reference(sheet, min_col=int(chart_spec.get("category_col", 1)), min_row=2, max_row=sheet.max_row)
                chart.add_data(data, titles_from_data=True); chart.set_categories(categories)
                chart.title = str(chart_spec.get("title", "Resumen")); sheet.add_chart(chart, str(chart_spec.get("anchor", "E2")))
            sheet.freeze_panes = "A2"
        workbook.save(spec.path)

    @staticmethod
    def _create_pptx(spec: ArtifactSpec) -> None:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt

        def add_native_bullet(paragraph) -> None:
            properties = paragraph._p.get_or_add_pPr()
            for child in list(properties):
                if child.tag.endswith(("buNone", "buChar", "buAutoNum")):
                    properties.remove(child)
            bullet = OxmlElement("a:buChar"); bullet.set("char", "•")
            properties.append(bullet)
            properties.set("marL", str(int(Inches(.42))))
            properties.set("indent", str(-int(Inches(.25))))

        def add_cover_picture(slide, image_path: Path, *, alt_text: str) -> None:
            """Place a picture in a frame without distorting its aspect ratio."""
            from PIL import Image

            frame_left, frame_top = Inches(7.25), Inches(0)
            frame_width, frame_height = Inches(6.083), Inches(7.5)
            with Image.open(image_path) as image:
                source_ratio = image.width / image.height
            frame_ratio = frame_width / frame_height
            picture = slide.shapes.add_picture(str(image_path), frame_left, frame_top, width=frame_width, height=frame_height)
            if source_ratio > frame_ratio:
                visible_fraction = frame_ratio / source_ratio
                crop = max(0.0, (1.0 - visible_fraction) / 2.0)
                picture.crop_left = crop; picture.crop_right = crop
            elif source_ratio < frame_ratio:
                visible_fraction = source_ratio / frame_ratio
                crop = max(0.0, (1.0 - visible_fraction) / 2.0)
                picture.crop_top = crop; picture.crop_bottom = crop
            picture._element.nvPicPr.cNvPr.set("descr", alt_text)

        def add_content_picture(slide, image_path: Path, *, alt_text: str, trim_bottom: float = 0.0) -> None:
            """Crop a picture to a stable 4.75 × 4.85 inch editorial frame."""
            from PIL import Image

            frame_left, frame_top = Inches(7.85), Inches(1.62)
            frame_width, frame_height = Inches(4.75), Inches(4.85)
            with Image.open(image_path) as image:
                source_ratio = image.width / image.height
            frame_ratio = frame_width / frame_height
            picture = slide.shapes.add_picture(str(image_path), frame_left, frame_top, width=frame_width, height=frame_height)
            if source_ratio > frame_ratio:
                visible_fraction = frame_ratio / source_ratio
                crop = max(0.0, (1.0 - visible_fraction) / 2.0)
                picture.crop_left = crop; picture.crop_right = crop
            elif source_ratio < frame_ratio:
                visible_fraction = source_ratio / frame_ratio
                crop = max(0.0, (1.0 - visible_fraction) / 2.0)
                picture.crop_top = crop; picture.crop_bottom = crop
            if trim_bottom:
                picture.crop_bottom = min(0.42, float(picture.crop_bottom) + max(0.0, min(0.30, trim_bottom)))
            picture._element.nvPicPr.cNvPr.set("descr", alt_text)
            outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, frame_left, frame_top, frame_width, frame_height)
            outline.fill.background(); outline.line.color.rgb = RGBColor(210, 222, 229); outline.line.width = Pt(1.25)

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]
        navy = RGBColor(20, 61, 89)
        teal = RGBColor(42, 157, 143)
        ink = RGBColor(28, 48, 65)
        paper = RGBColor(246, 248, 251)
        for index, slide_spec in enumerate(spec.slides or [{"title": spec.title, "body": spec.content}]):
            slide = presentation.slides.add_slide(blank)
            background = slide.background.fill
            background.solid(); background.fore_color.rgb = navy if index == 0 else paper
            title = str(slide_spec.get("title", ""))
            body = str(slide_spec.get("body", ""))
            image_value = str(slide_spec.get("image", "")).strip()
            image_path = Path(image_value).expanduser().resolve() if image_value else None
            if image_path is not None and not image_path.is_file():
                raise FileNotFoundError(f"presentation_image_missing:{image_path}")
            image_alt = str(slide_spec.get("image_alt") or f"Ilustración sobre {title}").strip()
            if index == 0:
                if image_path is not None:
                    add_cover_picture(slide, image_path, alt_text=image_alt)
                    veil = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.55), Inches(0), Inches(1.7), Inches(7.5))
                    veil.fill.solid(); veil.fill.fore_color.rgb = navy; veil.fill.transparency = 35; veil.line.fill.background()
                title_width = Inches(5.85) if image_path is not None else Inches(11.3)
                box = slide.shapes.add_textbox(Inches(.85 if image_path is not None else 1.0), Inches(2.0), title_width, Inches(1.55))
                frame = box.text_frame; frame.clear(); frame.word_wrap = True
                paragraph = frame.paragraphs[0]; paragraph.text = title; paragraph.alignment = PP_ALIGN.LEFT if image_path is not None else PP_ALIGN.CENTER
                paragraph.font.name = "Aptos Display"; paragraph.font.size = Pt(42 if image_path is not None else 44); paragraph.font.bold = True; paragraph.font.color.rgb = RGBColor(255, 255, 255)
                subtitle = slide.shapes.add_textbox(Inches(.88 if image_path is not None else 1.8), Inches(3.75), Inches(5.75 if image_path is not None else 9.7), Inches(1.0))
                sf = subtitle.text_frame; sf.clear(); sp = sf.paragraphs[0]; sp.text = body; sp.alignment = PP_ALIGN.LEFT if image_path is not None else PP_ALIGN.CENTER
                sp.font.name = "Aptos"; sp.font.size = Pt(22 if image_path is not None else 24); sp.font.color.rgb = RGBColor(210, 235, 236)
                label = slide.shapes.add_textbox(Inches(.9 if image_path is not None else 5.0), Inches(1.35), Inches(4.2 if image_path is not None else 3.3), Inches(.4))
                cover_label = str(spec.metadata.get("cover_label") or spec.title).upper()[:40]
                lp = label.text_frame.paragraphs[0]; lp.text = cover_label; lp.alignment = PP_ALIGN.LEFT if image_path is not None else PP_ALIGN.CENTER
                lp.font.name = "Aptos"; lp.font.size = Pt(13); lp.font.bold = True; lp.font.color.rgb = RGBColor(91, 214, 201)
            else:
                band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(.28), Inches(7.5))
                band.fill.solid(); band.fill.fore_color.rgb = teal; band.line.fill.background()
                title_box = slide.shapes.add_textbox(Inches(.85), Inches(.65), Inches(11.2), Inches(.85))
                tf = title_box.text_frame; tf.clear(); tp = tf.paragraphs[0]; tp.text = title
                tp.font.name = "Aptos Display"; tp.font.size = Pt(32); tp.font.bold = True; tp.font.color.rgb = navy
                if image_path is not None:
                    add_content_picture(
                        slide, image_path, alt_text=image_alt,
                        trim_bottom=float(slide_spec.get("image_trim_bottom", 0.0) or 0.0),
                    )
                body_width = Inches(6.25) if image_path is not None else Inches(10.9)
                body_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.85), body_width, Inches(4.7))
                bf = body_box.text_frame; bf.clear(); bf.word_wrap = True
                lines = [line.strip() for line in body.splitlines() if line.strip()] or [body]
                for line_index, line in enumerate(lines):
                    bp = bf.paragraphs[0] if line_index == 0 else bf.add_paragraph()
                    bp.text = line
                    if len(lines) > 1:
                        add_native_bullet(bp)
                    bp.font.name = "Aptos"; bp.font.size = Pt(24); bp.font.color.rgb = ink
                    bp.space_after = Pt(13); bp.level = 0
                number = slide.shapes.add_textbox(Inches(11.65), Inches(6.65), Inches(.8), Inches(.35))
                np = number.text_frame.paragraphs[0]; np.text = f"{index + 1:02d}"; np.alignment = PP_ALIGN.RIGHT
                np.font.name = "Aptos"; np.font.size = Pt(12); np.font.bold = True; np.font.color.rgb = teal
        presentation.core_properties.title = spec.title
        presentation.core_properties.comments = "ARCHI styled presentation v2"
        presentation.save(spec.path)

    @staticmethod
    def _create_pdf(spec: ArtifactSpec) -> None:
        _write_visual_pdf(spec.path, spec.title, spec.content, spec.sections)

    @staticmethod
    def _create_png(spec: ArtifactSpec) -> None:
        from PIL import Image, ImageDraw, ImageFont, PngImagePlugin

        width = max(720, min(int(spec.metadata.get("width", 1200)), 2400))
        height = max(900, min(int(spec.metadata.get("height", 1600)), 3000))
        image = Image.new("RGB", (width, height), "#F5F7FB")
        draw = ImageDraw.Draw(image)
        bold_candidates = (
            "C:/Windows/Fonts/seguisb.ttf", "C:/Windows/Fonts/arialbd.ttf",
        )
        regular_candidates = (
            "C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/arial.ttf",
        )
        def font(candidates: tuple[str, ...], size: int):
            for candidate in candidates:
                if Path(candidate).is_file():
                    return ImageFont.truetype(candidate, size)
            return ImageFont.load_default()
        title_font = font(bold_candidates, 58)
        item_font = font(bold_candidates, 34)
        body_font = font(regular_candidates, 25)
        draw.rounded_rectangle((45, 45, width - 45, 255), 34, fill="#143D59")
        draw.multiline_text((85, 82), spec.title, font=title_font, fill="white", spacing=10)
        items = list(spec.metadata.get("items", [])) or [
            {"title": section.get("heading", ""), "body": section.get("body", "")}
            for section in spec.sections
        ]
        palette = ("#2A9D8F", "#E9C46A", "#F4A261", "#E76F51", "#457B9D")
        top = 300
        available = height - top - 70
        card_height = max(150, (available - 24 * max(0, len(items) - 1)) // max(1, len(items)))
        for index, item in enumerate(items[:7]):
            y = top + index * (card_height + 24)
            draw.rounded_rectangle((65, y, width - 65, y + card_height), 26, fill="white", outline="#DDE4EC", width=3)
            draw.ellipse((90, y + 35, 160, y + 105), fill=palette[index % len(palette)])
            number_font = font(bold_candidates, 31)
            number_color = "#FFFFFF" if index in {0, 3, 4} else "#102A43"
            draw.text((125, y + 70), str(index + 1), font=number_font, fill=number_color, anchor="mm")
            draw.text((190, y + 29), str(item.get("title", "")), font=item_font, fill="#143D59")
            body = textwrap.fill(str(item.get("body", "")), width=68)
            draw.multiline_text((190, y + 77), body, font=body_font, fill="#334E68", spacing=7)
        info = PngImagePlugin.PngInfo()
        info.add_text("Title", spec.title)
        info.add_text("Content", "\n".join(f"{item.get('title', '')}: {item.get('body', '')}" for item in items))
        image.save(spec.path, format="PNG", optimize=True, pnginfo=info)

    @staticmethod
    def _inspect_txt(path: Path) -> dict[str, Any]:
        text = path.read_text(encoding="utf-8")
        return {"valid": True, "characters": len(text), "lines": len(text.splitlines()), "text": text[:100_000]}

    _inspect_md = _inspect_txt
    @staticmethod
    def _inspect_css(path: Path) -> dict[str, Any]:
        text = path.read_text(encoding="utf-8")
        valid = bool(text.strip()) and text.count("{") == text.count("}") and "@media" in text
        return {"valid": valid, "characters": len(text), "lines": len(text.splitlines()), "responsive": "@media" in text}

    @staticmethod
    def _inspect_js(path: Path) -> dict[str, Any]:
        text = path.read_text(encoding="utf-8")
        pairs = (("{", "}"), ("(", ")"), ("[", "]"))
        balanced = all(text.count(left) == text.count(right) for left, right in pairs)
        valid = bool(text.strip()) and balanced and "addEventListener" in text
        return {"valid": valid, "characters": len(text), "lines": len(text.splitlines()), "balanced_delimiters": balanced, "interactive": "addEventListener" in text}

    @staticmethod
    def _inspect_py(path: Path) -> dict[str, Any]:
        text = path.read_text(encoding="utf-8")
        ast.parse(text, filename=str(path))
        return {"valid": True, "characters": len(text), "lines": len(text.splitlines()), "syntax": "valid"}

    @staticmethod
    def _inspect_html(path: Path) -> dict[str, Any]:
        from html.parser import HTMLParser
        text = path.read_text(encoding="utf-8")
        parser = HTMLParser(); parser.feed(text); parser.close()
        references = re.findall(r'(?:src|href)=["\']([^"\']+)["\']', text, flags=re.I)
        missing = [
            ref for ref in references
            if not re.match(r"^(?:https?:|data:|#|mailto:)", ref, flags=re.I)
            and not (path.parent / ref.split("?", 1)[0].split("#", 1)[0]).is_file()
        ]
        return {
            "valid": bool(re.search(r"<!doctype\s+html", text, flags=re.I)) and not missing,
            "characters": len(text), "references": references, "missing_resources": missing,
        }

    @staticmethod
    def _inspect_png(path: Path) -> dict[str, Any]:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            return {
                "valid": image.format == "PNG" and image.width >= 64 and image.height >= 64,
                "width": image.width, "height": image.height,
                "title": str(image.info.get("Title", "")),
                "content": str(image.info.get("Content", "")),
            }

    @staticmethod
    def _inspect_json(path: Path) -> dict[str, Any]:
        data = json.loads(path.read_text(encoding="utf-8"))
        return {"valid": True, "type": type(data).__name__, "items": len(data) if hasattr(data, "__len__") else 1}

    @staticmethod
    def _inspect_csv(path: Path) -> dict[str, Any]:
        with path.open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source))
        return {"valid": True, "rows": len(rows), "columns": max((len(row) for row in rows), default=0)}

    @staticmethod
    def _inspect_docx(path: Path) -> dict[str, Any]:
        from docx import Document

        document = Document(path)
        return {"valid": True, "paragraphs": len(document.paragraphs), "tables": len(document.tables), "title": document.core_properties.title or ""}

    @staticmethod
    def _inspect_xlsx(path: Path) -> dict[str, Any]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets: dict[str, Any] = {}
            for sheet in workbook.worksheets:
                sample = [
                    [cell.value for cell in row[:20]]
                    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50))
                ]
                sheets[sheet.title] = {"rows": sheet.max_row, "columns": sheet.max_column, "sample": sample}
            with zipfile.ZipFile(path) as archive:
                charts = sum(name.startswith("xl/charts/chart") for name in archive.namelist())
                integrity = archive.testzip() is None
            return {"valid": bool(sheets) and integrity, "sheets": sheets, "sheet_count": len(sheets), "charts": charts}
        finally:
            workbook.close()

    @staticmethod
    def _inspect_pptx(path: Path) -> dict[str, Any]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(path)
        titles: list[str] = []
        text = []
        pictures_by_slide: list[int] = []
        for index, slide in enumerate(presentation.slides, 1):
            if index > 100:
                break
            values = [str(shape.text).strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and str(shape.text).strip()]
            native_title = str(slide.shapes.title.text).strip() if slide.shapes.title else ""
            titles.append(native_title or (values[0] if values else ""))
            text.append({"slide": index, "text": "\n".join(values)[:20_000]})
            pictures_by_slide.append(sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes))
        with zipfile.ZipFile(path) as archive:
            media_files = sorted(name for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/"))
            integrity = archive.testzip() is None
        return {
            "valid": bool(titles) and any(titles) and integrity, "slides": len(titles), "titles": titles, "content": text,
            "pictures": sum(pictures_by_slide), "pictures_by_slide": pictures_by_slide,
            "slides_with_pictures": sum(count > 0 for count in pictures_by_slide), "media_files": media_files,
        }

    @staticmethod
    def _inspect_pdf(path: Path) -> dict[str, Any]:
        from pypdf import PdfReader

        reader = PdfReader(path)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return {"valid": bool(reader.pages), "pages": len(reader.pages), "characters": len(text), "text": text[:100_000]}

    @staticmethod
    def _edit_txt(source: Path, output: Path, replacements: dict[str, str], append: str, _cells: dict[str, Any]) -> None:
        text = source.read_text(encoding="utf-8")
        for old, new in replacements.items():
            text = text.replace(old, new)
        if append:
            text += append
        output.write_text(text, encoding="utf-8")

    _edit_md = _edit_txt
    _edit_html = _edit_txt
    _edit_css = _edit_txt
    _edit_js = _edit_txt
    _edit_py = _edit_txt

    @staticmethod
    def _edit_json(source: Path, output: Path, replacements: dict[str, str], append: str, cells: dict[str, Any]) -> None:
        data = json.loads(source.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("json_edit_requires_object")
        data.update(cells)
        output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def _edit_csv(source: Path, output: Path, replacements: dict[str, str], append: str, cells: dict[str, Any]) -> None:
        text = source.read_text(encoding="utf-8-sig")
        for old, new in replacements.items():
            text = text.replace(old, new)
        output.write_text(text + append, encoding="utf-8-sig")

    @staticmethod
    def _edit_docx(source: Path, output: Path, replacements: dict[str, str], append: str, _cells: dict[str, Any]) -> None:
        from docx import Document

        document = Document(source)
        for paragraph in document.paragraphs:
            for old, new in replacements.items():
                if old in paragraph.text:
                    for run in paragraph.runs:
                        run.text = run.text.replace(old, new)
        if append:
            document.add_paragraph(append)
        document.save(output)

    @staticmethod
    def _edit_xlsx(source: Path, output: Path, replacements: dict[str, str], append: str, cells: dict[str, Any]) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(source)
        for reference, value in cells.items():
            if "!" not in reference:
                raise ValueError("xlsx_cell_reference_requires_sheet")
            sheet, cell = reference.split("!", 1); workbook[sheet][cell] = value
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        for old, new in replacements.items():
                            cell.value = cell.value.replace(old, new)
        workbook.save(output)

    @staticmethod
    def _edit_pptx(source: Path, output: Path, replacements: dict[str, str], append: str, _cells: dict[str, Any]) -> None:
        from pptx import Presentation

        presentation = Presentation(source)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            for old, new in replacements.items():
                                run.text = run.text.replace(old, new)
        if append:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Anexo"; slide.placeholders[1].text = append
        presentation.save(output)

    @staticmethod
    def _edit_pdf(source: Path, output: Path, replacements: dict[str, str], append: str, _cells: dict[str, Any]) -> None:
        if replacements:
            raise ValueError("pdf_text_replacement_not_lossless")
        from pypdf import PdfReader, PdfWriter

        writer = PdfWriter(); writer.append(source)
        if append:
            temporary = output.with_suffix(".append.pdf")
            _write_simple_pdf(temporary, append.splitlines())
            try:
                for page in PdfReader(temporary).pages:
                    writer.add_page(page)
            finally:
                temporary.unlink(missing_ok=True)
        with output.open("xb") as stream:
            writer.write(stream)


def _write_visual_pdf(path: Path, title: str, subtitle: str, sections: list[dict[str, Any]]) -> None:
    """Write a styled, text-extractable Letter PDF with deterministic pagination."""
    def escaped(value: str) -> str:
        raw = value.encode("cp1252", "replace").decode("latin-1")
        return raw.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    def text_command(value: str, x: float, y: float, size: float, *, bold: bool = False, color: str = "0.11 0.19 0.25") -> str:
        font = "F2" if bold else "F1"
        return f"BT /{font} {size:g} Tf {color} rg {x:g} {y:g} Td ({escaped(value)}) Tj ET"

    def page_header(page_number: int) -> tuple[list[str], float]:
        commands = [
            "0.965 0.973 0.984 rg 0 0 612 792 re f",
            "0.078 0.239 0.349 rg 0 650 612 142 re f",
            "0.165 0.616 0.561 rg 0 642 612 8 re f",
        ]
        title_lines = textwrap.wrap(title or "Documento PDF", width=38)[:2]
        title_y = 735
        for line in title_lines:
            commands.append(text_command(line, 48, title_y, 24, bold=True, color="1 1 1")); title_y -= 29
        sub_lines = textwrap.wrap(subtitle, width=74)[:2]
        sub_y = min(684, title_y - 8)
        for line in sub_lines:
            commands.append(text_command(line, 49, sub_y, 11, color="0.84 0.93 0.94")); sub_y -= 15
        commands.append(text_command(f"{page_number:02d}", 538, 32, 9, bold=True, color="0.165 0.616 0.561"))
        return commands, 612.0

    pages: list[list[str]] = []
    commands, y = page_header(1)
    for section_index, section in enumerate(sections):
        heading = str(section.get("heading", "")).strip()
        body = str(section.get("body", "")).strip()
        chunks = [item.strip() for item in body.split("|") if item.strip()] if "|" in body else [body]
        body_lines: list[str] = []
        for item in chunks:
            prefix = "- " if len(chunks) > 1 else ""
            wrapped = textwrap.wrap(prefix + item, width=82, subsequent_indent="  " if prefix else "") or [""]
            body_lines.extend(wrapped)
        for row in section.get("rows") or []:
            body_lines.extend(textwrap.wrap(" | ".join(str(value) for value in row), width=82) or [""])
        card_height = 44 + 14 * len(body_lines)
        if y - card_height < 58:
            pages.append(commands)
            commands, y = page_header(len(pages) + 1)
        fill = "0.925 0.956 0.957" if section_index % 2 == 0 else "1 1 1"
        commands.append(f"{fill} rg 42 {y-card_height:g} 528 {card_height:g} re f")
        commands.append(f"0.84 0.88 0.91 RG 0.7 w 42 {y-card_height:g} 528 {card_height:g} re S")
        commands.append(text_command(heading, 58, y - 25, 12, bold=True, color="0.078 0.239 0.349"))
        line_y = y - 47
        for line in body_lines:
            commands.append(text_command(line, 58, line_y, 10.5)); line_y -= 14
        y -= card_height + 12
    pages.append(commands)

    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"]
    page_ids: list[int] = []
    for page in pages:
        body = "\n".join(page).encode("latin-1")
        content_id = len(objects) + 1
        objects.append(f"<< /Length {len(body)} >>\nstream\n".encode() + body + b"\nendstream")
        page_id = len(objects) + 1; page_ids.append(page_id)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>".encode())
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    stream = io.BytesIO(); stream.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, obj in enumerate(objects, 1):
        offsets.append(stream.tell()); stream.write(f"{object_id} 0 obj\n".encode()); stream.write(obj); stream.write(b"\nendobj\n")
    xref = stream.tell(); stream.write(f"xref\n0 {len(objects)+1}\n".encode()); stream.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        stream.write(struct.pack(f">{10}s", str(offset).zfill(10).encode()) + b" 00000 n \n")
    stream.write(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(stream.getvalue())


def _write_simple_pdf(path: Path, lines: list[str]) -> None:
    """Write a compact WinAnsi text PDF without a resident rendering engine."""
    def escaped(value: str) -> str:
        raw = value.encode("cp1252", "replace").decode("latin-1")
        return raw.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    wrapped: list[str] = []
    for line in lines:
        wrapped.extend(textwrap.wrap(str(line), width=92, replace_whitespace=False, drop_whitespace=True) or [""])
    page_lines = [wrapped[index:index + 45] for index in range(0, max(1, len(wrapped)), 45)] or [[]]
    objects: list[bytes] = []
    page_ids: list[int] = []
    font_id = 3
    objects.extend([b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"])
    for page in page_lines:
        content = ["BT /F1 12 Tf 54 760 Td 15 TL"]
        for line in page:
            content.append(f"({escaped(str(line))}) Tj T*")
        content.append("ET")
        body = "\n".join(content).encode("latin-1")
        content_id = len(objects) + 1
        objects.append(f"<< /Length {len(body)} >>\nstream\n".encode() + body + b"\nendstream")
        page_id = len(objects) + 1; page_ids.append(page_id)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode())
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    stream = io.BytesIO(); stream.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, body in enumerate(objects, 1):
        offsets.append(stream.tell()); stream.write(f"{object_id} 0 obj\n".encode()); stream.write(body); stream.write(b"\nendobj\n")
    xref = stream.tell(); stream.write(f"xref\n0 {len(objects)+1}\n".encode()); stream.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        stream.write(struct.pack(f">{10}s", str(offset).zfill(10).encode()) + b" 00000 n \n")
    stream.write(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(stream.getvalue())
