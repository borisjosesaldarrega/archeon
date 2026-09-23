"""Checkpointed document copy editing with reopen verification."""

from __future__ import annotations

import json
import shutil
import time
from pathlib import Path
from typing import Any, Mapping, Sequence

from .reader import DocumentReader


class DocumentEditor:
    EDITABLE = {".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".html", ".htm", ".docx"}

    def __init__(self, reader: DocumentReader | None = None) -> None:
        self.reader = reader or DocumentReader()

    def edit_copy(
        self, source: str | Path, output: str | Path, *,
        replacements: Mapping[str, str] | None = None,
        append_paragraphs: Sequence[str] = (), remove_paragraphs: Sequence[int] = (),
    ) -> dict[str, Any]:
        source_path = Path(source).expanduser().resolve()
        output_path = Path(output).expanduser().resolve()
        if not source_path.is_file():
            raise FileNotFoundError("document_not_found")
        if source_path.suffix.casefold() not in self.EDITABLE:
            raise ValueError("document_format_is_read_only")
        if output_path == source_path:
            raise ValueError("edit_requires_output_copy")
        if output_path.exists():
            raise FileExistsError("output_document_already_exists")
        if not output_path.parent.is_dir():
            raise FileNotFoundError("output_folder_not_found")
        replacements = {str(key): str(value) for key, value in (replacements or {}).items()}
        if any(not key for key in replacements):
            raise ValueError("empty_replacement_source")
        checkpoint_dir = output_path.parent / ".archeon-checkpoints"
        checkpoint_dir.mkdir(exist_ok=True)
        checkpoint = checkpoint_dir / f"{source_path.stem}-{time.time_ns()}{source_path.suffix}.bak"
        shutil.copy2(source_path, checkpoint)
        try:
            if source_path.suffix.casefold() == ".docx":
                counts = self._edit_docx(
                    source_path, output_path, replacements,
                    append_paragraphs=append_paragraphs, remove_paragraphs=remove_paragraphs,
                )
            else:
                counts = self._edit_text(source_path, output_path, replacements, append_paragraphs)
            reopened = self.reader.read(output_path)
            missing = [value for value in replacements.values() if value and value not in reopened.text]
            if missing:
                output_path.unlink(missing_ok=True)
                raise RuntimeError("document_reopen_verification_failed")
        except BaseException:
            raise
        return {
            "source": str(source_path), "output": str(output_path),
            "checkpoint": str(checkpoint), "replacements": counts,
            "paragraphs_appended": len(append_paragraphs),
            "paragraphs_removed": len(remove_paragraphs),
            "reopened": True, "verified": True,
        }

    @staticmethod
    def _edit_text(
        source: Path, output: Path, replacements: Mapping[str, str], append_paragraphs: Sequence[str],
    ) -> dict[str, int]:
        text = source.read_text(encoding="utf-8-sig", errors="strict")
        counts: dict[str, int] = {}
        for old, new in replacements.items():
            counts[old] = text.count(old)
            text = text.replace(old, new)
        if append_paragraphs:
            text = text.rstrip() + "\n\n" + "\n\n".join(str(item) for item in append_paragraphs) + "\n"
        if source.suffix.casefold() == ".json":
            parsed = json.loads(text)
            text = json.dumps(parsed, ensure_ascii=False, indent=2) + "\n"
        output.write_text(text, encoding="utf-8")
        return counts

    @staticmethod
    def _edit_docx(
        source: Path, output: Path, replacements: Mapping[str, str], *,
        append_paragraphs: Sequence[str], remove_paragraphs: Sequence[int],
    ) -> dict[str, int]:
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("docx_provider_not_installed") from error
        document = Document(str(source))
        counts = {key: 0 for key in replacements}

        def replace_runs(paragraph: Any) -> None:
            for old, new in replacements.items():
                while old in paragraph.text:
                    full = paragraph.text
                    start = full.index(old)
                    end = start + len(old)
                    cursor = 0
                    inserted = False
                    for run in paragraph.runs:
                        run_start, run_end = cursor, cursor + len(run.text)
                        cursor = run_end
                        if run_end <= start or run_start >= end:
                            continue
                        local_start = max(0, start - run_start)
                        local_end = min(len(run.text), end - run_start)
                        replacement = new if not inserted else ""
                        run.text = run.text[:local_start] + replacement + run.text[local_end:]
                        inserted = True
                    counts[old] += 1

        for paragraph in document.paragraphs:
            replace_runs(paragraph)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replace_runs(paragraph)
        for index in sorted({int(item) for item in remove_paragraphs}, reverse=True):
            if index < 0 or index >= len(document.paragraphs):
                raise IndexError("paragraph_index_out_of_range")
            paragraph = document.paragraphs[index]
            paragraph._element.getparent().remove(paragraph._element)
        for value in append_paragraphs:
            document.add_paragraph(str(value))
        document.save(str(output))
        return counts
