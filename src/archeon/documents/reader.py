"""Bounded document extraction with format dependencies imported only on demand."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any


@dataclass(frozen=True, slots=True)
class DocumentPage:
    number: int
    text: str
    accessible_text: bool = True

    def public(self, *, max_characters: int = 200_000) -> dict[str, Any]:
        text = self.text[:max_characters]
        return {
            "number": self.number, "text": text,
            "characters": len(self.text), "truncated": len(text) < len(self.text),
            "accessible_text": self.accessible_text,
        }


@dataclass(frozen=True, slots=True)
class DocumentContent:
    path: Path
    format: str
    pages: tuple[DocumentPage, ...]
    metadata: dict[str, Any] = field(default_factory=dict)
    headings: tuple[dict[str, Any], ...] = ()
    tables: tuple[dict[str, Any], ...] = ()

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)

    def public(self, *, max_characters: int = 200_000) -> dict[str, Any]:
        return {
            "path": str(self.path), "format": self.format,
            "page_count": len(self.pages),
            "pages": [page.public(max_characters=max_characters) for page in self.pages],
            "metadata": dict(self.metadata), "headings": list(self.headings),
            "tables": list(self.tables),
            "visual_fallback_required": any(not page.accessible_text for page in self.pages),
        }


class _VisibleHTML(HTMLParser):
    _BREAKS = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0
        self.headings: list[dict[str, Any]] = []
        self._heading: tuple[int, list[str]] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.casefold()
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth += 1
        if not self.hidden_depth and tag in self._BREAKS:
            self.parts.append("\n")
        if not self.hidden_depth and re.fullmatch(r"h[1-6]", tag):
            self._heading = (int(tag[1]), [])

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if tag in {"script", "style", "noscript"} and self.hidden_depth:
            self.hidden_depth -= 1
        if self._heading and tag == f"h{self._heading[0]}":
            value = " ".join(self._heading[1]).strip()
            if value:
                self.headings.append({"level": self._heading[0], "text": value, "page": 1})
            self._heading = None
        if not self.hidden_depth and tag in self._BREAKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.hidden_depth:
            return
        value = data.strip()
        if value:
            self.parts.append(value + " ")
            if self._heading:
                self._heading[1].append(value)


class DocumentReader:
    SUPPORTED = {".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".html", ".htm", ".pdf", ".docx"}
    MAX_FILE_BYTES = 512 * 1024 * 1024
    MAX_EXTRACTED_CHARACTERS = 2_000_000
    MAX_PAGE_CHARACTERS = 500_000

    def read(
        self, path: str | Path, *, page: int | None = None,
        page_range: tuple[int, int] | None = None,
    ) -> DocumentContent:
        target = Path(path).expanduser().resolve()
        if not target.is_file():
            raise FileNotFoundError("document_not_found")
        suffix = target.suffix.casefold()
        if suffix not in self.SUPPORTED:
            raise ValueError("unsupported_document_type")
        size = target.stat().st_size
        if size > self.MAX_FILE_BYTES:
            raise ValueError("document_too_large")
        if page is not None and page_range is not None:
            raise ValueError("page_and_range_are_mutually_exclusive")
        if suffix == ".pdf":
            content = self._read_pdf(target, page=page, page_range=page_range)
        elif suffix == ".docx":
            content = self._read_docx(target)
        else:
            content = self._read_textual(target, suffix)
        if (page is not None or page_range is not None) and suffix != ".pdf":
            if page != 1:
                raise IndexError("document_page_out_of_range")
        return content

    def search(self, path: str | Path, query: str, *, limit: int = 50) -> dict[str, Any]:
        needle = query.casefold().strip()
        if not needle:
            raise ValueError("search_query_required")
        content = self.read(path)
        matches: list[dict[str, Any]] = []
        for page in content.pages:
            for line_number, line in enumerate(page.text.splitlines(), 1):
                offset = line.casefold().find(needle)
                if offset < 0:
                    continue
                matches.append({
                    "page": page.number, "line": line_number,
                    "context": line[max(0, offset - 80):offset + len(query) + 120].strip(),
                })
                if len(matches) >= max(1, min(limit, 200)):
                    break
            if len(matches) >= max(1, min(limit, 200)):
                break
        return {"path": str(content.path), "query": query, "matches": matches}

    def _read_textual(self, path: Path, suffix: str) -> DocumentContent:
        headings: list[dict[str, Any]] = []
        tables: list[dict[str, Any]] = []
        parts: list[str] = []
        characters = 0
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
            for line in source:
                if characters >= self.MAX_EXTRACTED_CHARACTERS:
                    break
                allowed = self.MAX_EXTRACTED_CHARACTERS - characters
                value = line[:allowed]
                parts.append(value)
                characters += len(value)
        raw = "".join(parts)
        file_truncated = path.stat().st_size > len(raw.encode("utf-8", errors="replace"))
        if suffix == ".json":
            if file_truncated:
                text = raw
            else:
                parsed = json.loads(raw)
                text = json.dumps(parsed, ensure_ascii=False, indent=2)
        elif suffix in {".csv", ".tsv"}:
            dialect = "excel-tab" if suffix == ".tsv" else "excel"
            row_count = 0
            max_columns = 0
            rendered: list[str] = []
            for row in csv.reader(io.StringIO(raw), dialect=dialect):
                row_count += 1
                max_columns = max(max_columns, len(row))
                rendered.append(" | ".join(cell for cell in row))
            text = "\n".join(rendered)
            tables.append({"index": 0, "rows": row_count, "columns": max_columns})
        elif suffix in {".html", ".htm"}:
            parser = _VisibleHTML()
            parser.feed(raw)
            text = re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()
            headings.extend(parser.headings)
        else:
            text = raw
            if suffix in {".md", ".markdown"}:
                for line in raw.splitlines():
                    match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
                    if match:
                        headings.append({"level": len(match.group(1)), "text": match.group(2), "page": 1})
        truncated = file_truncated or len(text) > self.MAX_EXTRACTED_CHARACTERS
        text = text[:self.MAX_EXTRACTED_CHARACTERS]
        return DocumentContent(
            path, suffix.lstrip("."), (DocumentPage(1, text),),
            {"bytes": path.stat().st_size, "extraction_truncated": truncated}, tuple(headings), tuple(tables),
        )

    def _read_pdf(
        self, path: Path, *, page: int | None,
        page_range: tuple[int, int] | None = None,
    ) -> DocumentContent:
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise RuntimeError("pdf_provider_not_installed") from error
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("encrypted_pdf_requires_password")
            except Exception as error:
                raise ValueError("encrypted_pdf_requires_password") from error
        page_count = len(reader.pages)
        if page is not None and (page < 1 or page > page_count):
            raise IndexError("document_page_out_of_range")
        if page_range is not None:
            start, end = page_range
            if start < 1 or end < start or end > page_count or end - start > 49:
                raise IndexError("document_page_range_out_of_range")
            indices = list(range(start - 1, end))
        else:
            indices = [page - 1] if page is not None else list(range(page_count))
        pages: list[DocumentPage] = []
        remaining = self.MAX_EXTRACTED_CHARACTERS
        extraction_truncated = False
        for index in indices:
            text = (reader.pages[index].extract_text() or "").strip()
            allowed = min(self.MAX_PAGE_CHARACTERS, remaining)
            if len(text) > allowed:
                extraction_truncated = True
            text = text[:allowed]
            remaining -= len(text)
            pages.append(DocumentPage(index + 1, text, accessible_text=bool(text)))
            if remaining <= 0:
                break
        metadata = {
            "page_count": page_count, "bytes": path.stat().st_size,
            "extraction_truncated": extraction_truncated or len(pages) < len(indices),
        }
        if reader.metadata:
            metadata.update({
                "title": str(reader.metadata.title or ""),
                "author": str(reader.metadata.author or ""),
            })
        return DocumentContent(path, "pdf", tuple(pages), metadata)

    def _read_docx(self, path: Path) -> DocumentContent:
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("docx_provider_not_installed") from error
        document = Document(str(path))
        parts: list[str] = []
        headings: list[dict[str, Any]] = []
        for paragraph in document.paragraphs:
            if paragraph.text:
                parts.append(paragraph.text)
            if paragraph.style and paragraph.style.name.casefold().startswith("heading"):
                match = re.search(r"(\d+)$", paragraph.style.name)
                headings.append({
                    "level": int(match.group(1)) if match else 1,
                    "text": paragraph.text, "page": 1,
                })
        tables: list[dict[str, Any]] = []
        for index, table in enumerate(document.tables):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            tables.append({"index": index, "rows": rows, "columns": max((len(row) for row in rows), default=0)})
            parts.extend(" | ".join(row) for row in rows)
        text = "\n".join(parts)
        truncated = len(text) > self.MAX_EXTRACTED_CHARACTERS
        return DocumentContent(
            path, "docx", (DocumentPage(1, text[:self.MAX_EXTRACTED_CHARACTERS]),),
            {"bytes": path.stat().st_size, "paragraph_count": len(document.paragraphs), "table_count": len(document.tables), "extraction_truncated": truncated},
            tuple(headings), tuple(tables),
        )
