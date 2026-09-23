"""Command-line entry point for the runnable ARCHEON Core."""

from __future__ import annotations

import argparse
import ctypes
import json
import multiprocessing
import time
import webbrowser
from pathlib import Path
from threading import Event

from archeon.app import ArcheonApplication
from archeon.core.tools import ToolContext
from archeon.ui import DesktopHost, DesktopUnavailable


class SingleInstance:
    """Windows named mutex that prevents duplicated ARCHEON runtimes."""

    def __init__(self) -> None:
        self._handle = None

    def acquire(self) -> bool:
        if __import__("os").name != "nt":
            return True
        kernel32 = ctypes.windll.kernel32
        handle = kernel32.CreateMutexW(None, False, "Local\\ARCHEON.Core.SingleInstance")
        if not handle:
            return False
        if kernel32.GetLastError() == 183:
            kernel32.CloseHandle(handle)
            return False
        self._handle = handle
        return True

    def close(self) -> None:
        if self._handle:
            ctypes.windll.kernel32.CloseHandle(self._handle)
            self._handle = None


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="archeon")
    parser.add_argument("--headless", action="store_true", help="Run Core and UI server without a window")
    parser.add_argument("--browser", action="store_true", help="Developer diagnostics only: open the local UI in a browser")
    parser.add_argument("--ghost", action="store_true", help="Start with the independent Ghost window")
    parser.add_argument("--auto-exit", type=float, default=None, metavar="SECONDS")
    parser.add_argument("--data-dir", type=Path, default=None)
    parser.add_argument("--port", type=int, default=0)
    parser.add_argument("--console-log", action="store_true")
    parser.add_argument("--benchmark-music", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-guest", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-launcher", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-radial", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-background", choices=("image", "video"), help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-stt-worker", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-desktop-observe", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-desktop-documents", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--benchmark-artifact-package", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--plugin-host", type=Path, default=None, help=argparse.SUPPRESS)
    parser.add_argument("--plugin-ready-file", type=Path, default=None, help=argparse.SUPPRESS)
    parser.add_argument("--plugin-ready-token", default="", help=argparse.SUPPRESS)
    parser.add_argument("--allow-multiple", action="store_true", help=argparse.SUPPRESS)
    return parser


def main(argv: list[str] | None = None) -> int:
    multiprocessing.freeze_support()
    args = build_parser().parse_args(argv)
    if args.plugin_host is not None:
        from archeon.plugins.host import main as run_plugin_host
        return run_plugin_host(
            args.plugin_host, ready_file=args.plugin_ready_file,
            ready_token=args.plugin_ready_token,
        )
    instance = SingleInstance()
    if not args.allow_multiple and not instance.acquire():
        return 0
    application = ArcheonApplication(
        data_dir=args.data_dir,
        port=args.port,
        console_log=args.console_log,
    )
    application.start()
    if args.benchmark_desktop_observe:
        observed = application.handle_command("Archeon, mira la ventana y dime qué ves")
        evidence = observed.get("data", {}).get("evidence", {})
        (application.data_dir / "desktop-agent-smoke.json").write_text(json.dumps({
            "ok": observed.get("ok", False),
            "route": observed.get("data", {}).get("route"),
            "mode": observed.get("data", {}).get("mode"),
            "status": observed.get("data", {}).get("status"),
            "provider": evidence.get("provider"),
            "window_process": evidence.get("window", {}).get("process_name"),
            "element_count": evidence.get("element_count", 0),
            "state_hash": evidence.get("state_hash"),
        }, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.benchmark_desktop_documents:
        from docx import Document

        fixture = application.data_dir / "packaged-document.docx"
        document = Document()
        document.add_heading("ARCHEON Desktop Agent", level=1)
        document.add_paragraph("Packaged document reader verification.")
        document.save(fixture)
        context = ToolContext(
            "packaged-documents", scope_permissions=frozenset({"filesystem.read"}),
        )
        document_result = application.tools.execute(
            "documents.read", {"path": str(fixture)}, context=context,
        )
        project_result = application.tools.execute(
            "programming.detect", {"path": str(Path.cwd())}, context=context,
        )
        manifests = {item.id for item in application.tools.manifests()}
        report = {
            "ok": bool(document_result.ok and project_result.ok),
            "document_verified": bool(document_result.verified),
            "document_text_found": any(
                "Packaged document reader verification." in str(page.get("text", ""))
                for page in document_result.data.get("pages", [])
            ),
            "programming_detect_verified": bool(project_result.verified),
            "browser_registered": "browser.navigate" in manifests and "browser.read" in manifests,
            "vision_model_downloaded": False,
        }
        report["ok"] = bool(report["ok"] and report["document_text_found"] and report["browser_registered"])
        (application.data_dir / "desktop-agent-documents-smoke.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8",
        )
    if args.benchmark_artifact_package:
        fixture = application.data_dir / "packaged-artifacts"
        fixture.mkdir(exist_ok=True)
        loaded_tools_before = application.tools.loaded_tool_count
        context = ToolContext(
            "packaged-artifacts", scope_permissions=frozenset({"filesystem.read", "filesystem.write"}),
        )
        docx = application.tools.execute("artifacts.create", {
            "path": str(fixture / "report.docx"), "title": "Verification", "content": "Packaged artifact verification",
        }, context=context)
        xlsx = application.tools.execute("artifacts.create", {
            "path": str(fixture / "summary.xlsx"), "sheets": [{"name": "Data", "rows": [["Item", "Value"], ["Artifacts", 4]], "chart": {"title": "Artifacts"}}],
        }, context=context)
        pptx = application.tools.execute("artifacts.create", {
            "path": str(fixture / "brief.pptx"), "slides": [{"title": f"Slide {index}", "body": "Artifact verification"} for index in range(1, 7)],
        }, context=context)
        archive = application.tools.execute("archives.create", {
            "destination": str(fixture / "bundle.zip"), "sources": [docx.data.get("path", ""), xlsx.data.get("path", ""), pptx.data.get("path", "")],
        }, context=context)
        report = {
            "ok": all(item.ok and item.verified for item in (docx, xlsx, pptx, archive)),
            "startup_ms": round(application.startup_ms, 3),
            "loaded_tools_before": loaded_tools_before,
            "loaded_tools_after": application.tools.loaded_tool_count,
            "docx": {"verified": docx.verified, "bytes": docx.data.get("bytes")},
            "xlsx": {"verified": xlsx.verified, "charts": xlsx.data.get("structure", {}).get("charts")},
            "pptx": {"verified": pptx.verified, "slides": pptx.data.get("structure", {}).get("slides")},
            "archive": {"verified": archive.verified, "members": archive.data.get("member_count"), "sha256": archive.data.get("sha256")},
        }
        (application.data_dir / "artifact-package-smoke.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8",
        )
    if args.benchmark_stt_worker:
        started = time.perf_counter()
        provider = application.voice._stt  # Deliberately private: diagnostic-only path.
        provider.transcribe(bytes(32_000), 16_000)
        (application.data_dir / "stt-worker-smoke.json").write_text(json.dumps({
            "ok": True,
            "elapsed_ms": round((time.perf_counter() - started) * 1_000, 3),
            "model_loaded_after": provider.loaded,
        }), encoding="utf-8")
    if args.benchmark_background:
        visual = (
            application.resources.ui("logo_asitente.png")
            if args.benchmark_background == "image"
            else application.resources.asset("ARCHEON.mp4")
        )
        application.configuration.update_settings({"appearance": {
            "background_type": args.benchmark_background,
            "background_path": str(visual),
        }})
    launcher_ms = None
    launcher_count = None
    if args.benchmark_launcher:
        launcher_started = time.perf_counter()
        launcher_count = len(application.launcher.refresh(force=True))
        launcher_ms = (time.perf_counter() - launcher_started) * 1000
    print(
        json.dumps(
            {
                "event": "archeon.ready",
                "url": application.ui_server.url,
                "startup_ms": round(application.startup_ms, 3),
                "pid": __import__("os").getpid(),
            },
            separators=(",", ":"),
        ),
        flush=True,
    )
    if launcher_ms is not None:
        print(json.dumps({
            "event": "archeon.benchmark.launcher", "items": launcher_count,
            "scan_ms": round(launcher_ms, 3),
        }, separators=(",", ":")), flush=True)
    try:
        if args.headless:
            stopped = Event()
            stopped.wait(args.auto_exit)
        elif args.browser:
            webbrowser.open(f"{application.ui_server.url}/?token={application.ui_server.token}")
            if args.auto_exit:
                time.sleep(args.auto_exit)
            else:
                Event().wait()
        else:
            if args.benchmark_music and args.ghost:
                benchmark_result = application.handle_action("media.play")
                print(
                    json.dumps(
                        {"event": "archeon.benchmark.music", "ok": benchmark_result.get("ok", False)},
                        separators=(",", ":"),
                    ),
                    flush=True,
                )
            host = DesktopHost(
                application.events,
                base_url=application.ui_server.url,
                token=application.ui_server.token,
                config=application.configuration.config,
                action_handler=application.handle_action,
                media_status_handler=application.media.status,
                artwork_handler=application.media.artwork,
                resources=application.resources,
            )
            start_in_ghost = bool(
                args.ghost or application.configuration.config.ghost.enabled
                or application.configuration.config.startup.start_in_ghost_mode
            )
            host.run(
                initial_mode="ghost" if start_in_ghost else "main",
                auto_exit_seconds=args.auto_exit,
                benchmark_music=args.benchmark_music,
                benchmark_guest=args.benchmark_guest,
                benchmark_radial=args.benchmark_radial,
            )
    except KeyboardInterrupt:
        pass
    except DesktopUnavailable as error:
        print(json.dumps({"event": "archeon.error", "error": str(error)}), flush=True)
        return 2
    finally:
        application.stop()
        instance.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
