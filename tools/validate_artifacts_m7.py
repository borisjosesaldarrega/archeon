"""Reopen and independently validate every final M7 artifact."""

from __future__ import annotations

import csv
import json
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = Path.home() / "Downloads" / "ARCHI-Artifactos-M7"
sys.path.insert(0, str(ROOT / "src"))

from archeon.documents.quality import RequirementChecker  # noqa: E402


class _HTMLCheck(HTMLParser):
    def __init__(self) -> None:
        super().__init__(); self.tags: set[str] = set()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.tags.add(tag)


def _text_formats() -> dict[str, object]:
    txt = (OUT / "resumen_m7.txt").read_text(encoding="utf-8")
    md = (OUT / "README_M7.md").read_text(encoding="utf-8")
    html = (OUT / "M7_REPORT.html").read_text(encoding="utf-8")
    parser = _HTMLCheck(); parser.feed(html)
    with (OUT / "gastos_m7.csv").open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    json_document = json.loads((OUT / "gastos_m7.json").read_text(encoding="utf-8"))
    json_rows = json_document["gastos"]
    return {
        "txt": {"bytes": len(txt.encode("utf-8")), "utf8": True, "passed": bool(txt.strip())},
        "md": {"headings": md.count("\n#") + int(md.startswith("#")), "lists": "\n- " in md, "table": "|" in md, "links": "](" in md, "code_fence": "```" in md, "passed": all(("\n- " in md, "|" in md, "](" in md, "```" in md))},
        "html": {"tags": sorted(parser.tags), "passed": {"html", "head", "body", "main"}.issubset(parser.tags)},
        "csv": {"rows": len(csv_rows), "columns": list(csv_rows[0]) if csv_rows else [], "passed": bool(csv_rows)},
        "json": {"rows": len(json_rows), "parsed": isinstance(json_document, dict), "passed": isinstance(json_rows, list) and bool(json_rows)},
        "_csv_rows": csv_rows,
        "_json_rows": json_rows,
    }


def _xlsx(csv_rows: list[dict[str, str]], json_rows: list[dict[str, object]]) -> dict[str, object]:
    path = OUT / "Analisis_Gastos_Mensuales_ARCHI.xlsx"
    checker = RequirementChecker().check_xlsx(path, required_sheets=("Datos", "Resumen"), minimum_charts=1)
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        sheet = workbook["Datos"]
        months = [sheet.cell(row=row, column=1).value for row in range(4, 10)]
        values = [[sheet.cell(row=row, column=col).value for col in range(2, 8)] for row in range(4, 10)]
        formulas = [sheet.cell(row=row, column=8).value for row in range(4, 10)]
        formula_ok = formulas == [f"=SUM(B{row}:G{row})" for row in range(4, 10)]
        numeric_totals = [round(sum(float(value) for value in row), 2) for row in values]
        csv_keys = ("Alimentación", "Transporte", "Servicios", "Educación", "Salud", "Otros")
        json_keys = ("alimentacion", "transporte", "servicios", "educacion", "salud", "otros")
        csv_totals = [round(sum(float(row[key]) for key in csv_keys), 2) for row in csv_rows]
        json_totals = [round(sum(float(row[key]) for key in json_keys), 2) for row in json_rows]
        coherence = months == [row["Mes"] for row in csv_rows] == [row["mes"] for row in json_rows] and numeric_totals == csv_totals == json_totals
        checker.update({"monthly_totals": numeric_totals, "formula_pattern": formula_ok, "csv_json_coherence": coherence})
        checker["passed"] = bool(checker["passed"] and formula_ok and coherence and checker["tables"] >= 2)
        return checker
    finally:
        workbook.close()


def main() -> int:
    checker = RequirementChecker(); text_checks = _text_formats()
    csv_rows = text_checks.pop("_csv_rows"); json_rows = text_checks.pop("_json_rows")
    security_renders = sorted((ROOT / "tmp" / "m7-renders" / "pdf-security").glob("page-*.png"))
    info_renders = sorted((ROOT / "tmp" / "m7-renders" / "pdf-info").glob("page-*.png"))
    node = Path(r"C:\Users\salda\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
    js_check = subprocess.run([str(node), "--check", str(OUT / "ARCHI_TaskBoard" / "app.js")], capture_output=True, text=True, check=False)
    archive_validation = json.loads((OUT / "M7_ARCHIVE_VALIDATION.json").read_text(encoding="utf-8"))
    archives = {suffix: checker.check_archive(OUT / name, expected_files=("ARCHI_TaskBoard/index.html", "ARCHI_TaskBoard/styles.css", "ARCHI_TaskBoard/app.js", "ARCHI_TaskBoard/README.md")) for suffix, name in {"zip": "ARCHI_TaskBoard.zip", "tar": "ARCHI_TaskBoard.tar", "tar.gz": "ARCHI_TaskBoard.tar.gz"}.items()}
    report: dict[str, object] = {
        "milestone": "Artifact Intelligence M7",
        "xlsx": _xlsx(csv_rows, json_rows),
        "pptx": checker.check_pptx(OUT / "Presentacion_Seguridad_Web_ARCHI.pptx", expected_slides=7, rendered_slides=security_renders),
        "infographic": checker.check_infographic(OUT / "Infografia_Embriologia_ARCHI.pptx", required_phases=("Fecundación", "Segmentación", "Implantación", "Gastrulación", "Neurulación", "Organogénesis", "Periodo fetal"), rendered_slides=info_renders),
        "pdf": {"security_pages": len(PdfReader(OUT / "Presentacion_Seguridad_Web_ARCHI.pdf").pages), "infographic_pages": len(PdfReader(OUT / "Infografia_Embriologia_ARCHI.pdf").pages)},
        "web": {**checker.check_web_project(OUT / "ARCHI_TaskBoard"), "javascript_syntax": js_check.returncode == 0, "browser_agent_smoke": {"status": "BLOCKED_BY_BROWSER_URL_POLICY", "passed": False, "reason": "The authorized Browser tool refused local file navigation and explicitly prohibited workarounds."}},
        "archives": {"formats": archives, "native_creation_validation": archive_validation, "optional": {"7z": "UNAVAILABLE_OPTIONAL_PROVIDER", "rar": "UNAVAILABLE_LICENSED_PROVIDER"}},
        **text_checks,
    }
    pdf = report["pdf"]; pdf["passed"] = pdf["security_pages"] == 7 and pdf["infographic_pages"] == 1
    web = report["web"]; web["static_passed"] = bool(web["passed"] and web["javascript_syntax"]); web["passed"] = False
    archives_section = report["archives"]; archives_section["passed"] = all(value["passed"] for value in archives.values()) and archive_validation["all_native_passed"]
    required = ("xlsx", "pptx", "infographic", "pdf", "archives", "txt", "md", "csv", "json", "html")
    report["artifact_checks_passed"] = all(bool(report[name]["passed"]) for name in required)
    report["overall_status"] = "PARTIAL_BROWSER_AGENT_BLOCKED" if report["artifact_checks_passed"] else "FAILED"
    target = OUT / "M7_ARTIFACT_VALIDATION.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"path": str(target), "artifact_checks_passed": report["artifact_checks_passed"], "overall_status": report["overall_status"]}, ensure_ascii=False))
    return 0 if report["artifact_checks_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
