"""Build the evidence-led M9 multilingual report without loading AI/STT at idle."""

from __future__ import annotations

import json
import statistics
import sys
import tempfile
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from archeon.core.config import SUPPORTED_LOCALES
from archeon.core.language import LanguageContextEngine
from archeon.core.multilingual import LANGUAGE_PROBES, MultilingualQualityGate
from archeon.core.paths import AppPaths
from archeon.documents import DocumentStyleProfile
from archeon.ui.server import UI_ROOT
from archeon.voice.providers import WindowsSapiProvider


DOCUMENT_TEXT = {
    "es": "Documento de prueba en español.", "en": "English test document.",
    "pt": "Documento de teste em português.", "fr": "Document de test en français.",
    "de": "Testdokument auf Deutsch.", "it": "Documento di prova in italiano.",
    "zh": "中文测试文档。", "ja": "日本語のテスト文書です。", "ko": "한국어 테스트 문서입니다.",
    "ru": "Тестовый документ на русском языке.", "ar": "مستند اختبار باللغة العربية.",
    "hi": "हिन्दी में परीक्षण दस्तावेज़।",
}


def percentile(values: list[float], ratio: float) -> float:
    ordered = sorted(values)
    return ordered[min(len(ordered) - 1, round((len(ordered) - 1) * ratio))]


def main() -> int:
    paths = AppPaths.discover()
    voices = WindowsSapiProvider.voices()
    tts_locales = sorted({locale for voice in voices for locale in voice.get("locales", [])})
    gate = MultilingualQualityGate(UI_ROOT / "locales", paths.model_dir(None))
    report = gate.evaluate(tts_locales=tts_locales)

    engine = LanguageContextEngine()
    samples: list[float] = []
    for _ in range(1_000):
        for probe in LANGUAGE_PROBES.values():
            started = time.perf_counter()
            engine.detect(probe)
            samples.append((time.perf_counter() - started) * 1_000)
    report["language_detection_benchmark_ms"] = {
        "iterations": len(samples), "median": round(statistics.median(samples), 6),
        "p95": round(percentile(samples, 0.95), 6), "max": round(max(samples), 6),
    }

    conversation = []
    previous = None
    for text in (
        "Explícame DNS en español.", "Now summarize that in English.",
        "Agora crie um PDF em português.", "والآن اختصره بالعربية.",
    ):
        decision = engine.decide(text, preferred="auto", previous=previous, fallback="es")
        conversation.append({"input": text, "response_language": decision.response_language, "source": decision.source})
        previous = decision.response_language
    report["mixed_language_conversation"] = conversation

    from docx import Document
    document_results = []
    with tempfile.TemporaryDirectory(prefix="archeon-m9-languages-") as directory:
        root = Path(directory)
        for locale in SUPPORTED_LOCALES:
            output = root / f"document-{locale}.docx"
            profile = DocumentStyleProfile.for_locale(locale)
            document = Document(); profile.apply_docx(document); document.add_paragraph(DOCUMENT_TEXT[locale]); document.save(output)
            reopened = Document(output)
            recovered = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
            with zipfile.ZipFile(output) as archive:
                styles = archive.read("word/styles.xml").decode("utf-8")
            document_results.append({
                "locale": locale, "unicode_roundtrip": DOCUMENT_TEXT[locale] in recovered,
                "font": profile.font_family, "direction": profile.direction,
                "rtl_style": "w:bidi" in styles if locale == "ar" else None,
            })
    report["document_docx_fixtures"] = document_results
    report["tts_inventory"] = {
        "voice_count": len(voices), "runtime_locales": tts_locales,
        "voices": [{"name": item.get("name"), "locales": item.get("locales", [])} for item in voices],
        "audio_playback_tested": False,
    }
    report["verification"] = {
        "automated": True, "packaged": False, "native_speaker_review": False,
        "user_verified": False,
    }
    output = ROOT / "benchmarks" / "M9_MULTILINGUAL_QUALITY.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "output": str(output), "tts_locales": tts_locales, "locales": len(report["locales"])}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
