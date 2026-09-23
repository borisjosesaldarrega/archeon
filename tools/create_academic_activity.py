"""Create and structurally validate the real ARCHI academic document fixture."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


TITLE = "Actividad: Planificación y Optimización de una Página Web"
QUESTIONS = [
    ("Pregunta 1", "¿Qué aspectos se deben definir antes de comenzar a desarrollar una página web?", [
        "Antes de escribir código conviene aclarar para qué se hará la página, a quién estará dirigida y qué acción se espera de sus visitantes. No es igual diseñar una tienda, un portafolio o una página informativa. El objetivo ayuda a decidir qué contenido es necesario, qué tono usar y cuáles funciones realmente aportan valor.",
        "También se deben definir las secciones principales, los recursos disponibles, el tipo de dispositivo más usado por el público y los requisitos básicos de seguridad y accesibilidad. Es útil fijar un alcance realista, responsables y tiempos de revisión. Con esas decisiones se evita empezar con muchas ideas sueltas y terminar con una página difícil de mantener.",
    ]),
    ("Pregunta 2", "¿Por qué es importante organizar previamente la estructura y el contenido de un sitio web?", [
        "Organizar primero la estructura permite que cada página tenga una función clara y que el usuario encuentre la información siguiendo un recorrido lógico. Un mapa sencillo del sitio y un esquema de cada pantalla muestran dónde irán el menú, los títulos, las llamadas a la acción y los contenidos secundarios antes de invertir tiempo en detalles visuales.",
        "Preparar el contenido con anticipación también evita diseñar espacios que luego no se ajustan al texto o a las imágenes reales. Además, facilita detectar repeticiones, vacíos y páginas innecesarias. El resultado suele ser más coherente y requiere menos correcciones porque diseño, contenido y desarrollo avanzan con la misma referencia.",
    ]),
    ("Pregunta 3", "¿Qué factores deben tomarse en cuenta para que una página web sea fácil de usar?", [
        "La navegación debe ser comprensible, con nombres de secciones claros y controles que se comporten como el usuario espera. El texto necesita buen contraste, tamaño legible y párrafos fáciles de recorrer. Los formularios deben indicar qué información solicitan, explicar los errores y permitir completar la tarea sin pasos innecesarios.",
        "También importa que la página responda bien en pantallas pequeñas, pueda usarse con teclado y muestre una respuesta visible después de cada acción. La velocidad forma parte de la facilidad de uso: si una pantalla tarda demasiado o cambia de posición mientras carga, la experiencia se vuelve confusa. Por eso conviene probar con personas, dispositivos y conexiones diferentes, no solo en la computadora donde se desarrolló.",
    ]),
    ("Pregunta 4", "Menciona cinco acciones que pueden ayudar a optimizar el rendimiento de una página web y explica brevemente cada una.", [
        "Optimizar imágenes: elegir el formato adecuado, reducir dimensiones y comprimirlas antes de publicarlas disminuye mucho la transferencia. Las imágenes que están fuera de la vista inicial pueden cargarse cuando el usuario se acerque a ellas.",
        "Reducir JavaScript: retirar bibliotecas y funciones que no se usan evita descargar, interpretar y ejecutar código innecesario. Los módulos secundarios pueden cargarse bajo demanda.",
        "Optimizar CSS: eliminar reglas duplicadas, dividir estilos críticos de los secundarios y comprimir el archivo reduce trabajo de red y de renderizado sin cambiar el diseño.",
        "Disminuir solicitudes: agrupar recursos pequeños cuando sea razonable y evitar fuentes, iconos o rastreadores redundantes reduce viajes entre el navegador y el servidor.",
        "Configurar caché: permitir que recursos versionados se reutilicen en visitas posteriores evita descargar lo mismo cada vez. Cuando un archivo cambia, su nombre o versión debe cambiar para que el navegador reciba la actualización.",
    ]),
    ("Pregunta 5", "¿Qué problemas pueden aparecer si una página web se desarrolla sin planificación ni pruebas previas?", [
        "Sin planificación pueden aparecer secciones repetidas, navegación desordenada, funciones que no apoyan el objetivo y cambios constantes de alcance. El equipo puede rehacer pantallas completas porque descubrió tarde que faltaba contenido, que el flujo no era comprensible o que una decisión técnica impedía crecer.",
        "Sin pruebas también pueden quedar enlaces rotos, errores en formularios, incompatibilidades con móviles, problemas de teclado o contraste y tiempos de carga demasiado altos. Estos fallos afectan la confianza del usuario y encarecen el mantenimiento. Probar antes de publicar no elimina todos los riesgos, pero permite corregir los más visibles con menor costo.",
    ]),
]

TABLE_ROWS = [
    ["Imágenes pesadas", "La página tarda en mostrar el contenido y consume más datos.", "Redimensionar, comprimir y usar formatos adecuados; aplicar carga diferida cuando corresponda."],
    ["JavaScript innecesario", "Aumenta el tiempo de descarga y bloquea el hilo principal.", "Eliminar dependencias sin uso y cargar módulos secundarios bajo demanda."],
    ["CSS sin optimizar", "Se transfieren reglas repetidas y el navegador procesa estilos que no necesita.", "Depurar reglas, separar estilos críticos y comprimir el archivo final."],
    ["Muchas solicitudes HTTP", "Cada recurso agrega espera y puede retrasar la vista inicial.", "Eliminar recursos redundantes y agrupar solo cuando reduzca solicitudes sin perjudicar la caché."],
    ["Falta de caché", "Los visitantes descargan de nuevo recursos que no cambiaron.", "Definir políticas de caché para archivos versionados y actualizar su versión al modificarlos."],
]

PRACTICES = [
    "Definir objetivos, público y alcance antes de diseñar las pantallas.",
    "Crear un mapa del sitio y revisar los recorridos principales del usuario.",
    "Diseñar primero con contenido real o con muestras cercanas a su extensión final.",
    "Medir el rendimiento en móvil y con una conexión limitada, no solo en una PC rápida.",
    "Comprobar navegación por teclado, contraste, textos alternativos y etiquetas de formularios.",
    "Mantener dependencias, estilos y scripts bajo control, retirando lo que ya no se utiliza.",
]


def _set_font(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)


def _fixed_table(table, widths: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    width = width if width is not None else OxmlElement("w:tblW")
    width.set(qn("w:type"), "dxa"); width.set(qn("w:w"), str(sum(widths)))
    if width.getparent() is None: properties.append(width)
    indent = properties.find(qn("w:tblInd"))
    indent = indent if indent is not None else OxmlElement("w:tblInd")
    indent.set(qn("w:type"), "dxa"); indent.set(qn("w:w"), "120")
    if indent.getparent() is None: properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol"); column.set(qn("w:w"), str(value)); grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths):
            cell.width = Inches(value / 1440)
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW(); tcw.type = "dxa"; tcw.w = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar"); cell._tc.get_or_add_tcPr().append(margins)
            for side, amount in (("top", 100), ("bottom", 100), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}")
                node.set(qn("w:w"), str(amount)); node.set(qn("w:type"), "dxa")
                if node.getparent() is None: margins.append(node)
    header = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); header.append(repeat)


def build(output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists(): output.unlink()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after, color in (("Heading 1",16,16,8,"2E74B5"),("Heading 2",13,12,6,"2E74B5"),("Heading 3",12,8,4,"1F4D78")):
        style = styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
    title = document.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    _set_font(title.add_run(TITLE), 20, bold=True, color="1F4D78")
    intro_heading = document.add_paragraph("Introducción", style="Heading 1")
    document.add_paragraph(
        "Planificar una página web permite transformar una idea general en un sitio útil, ordenado y posible de mantener. La optimización no comienza al final: está presente cuando se decide qué contenido mostrar, cómo organizarlo y qué recursos necesita cada pantalla. En esta actividad se revisan decisiones prácticas que ayudan a construir una web clara, rápida y fácil de probar."
    )
    for label, question, answers in QUESTIONS:
        document.add_paragraph(label, style="Heading 1")
        q = document.add_paragraph(); q.paragraph_format.keep_with_next = True
        run = q.add_run(question); _set_font(run, 11, bold=True)
        for answer in answers:
            document.add_paragraph(answer)
    document.add_paragraph("Problemas frecuentes de optimización", style="Heading 1")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Problema", "Consecuencia", "Posible solución"]):
        cell.text = value
        shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E8EEF5"); cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs: _set_font(run, 10, bold=True, color="1F4D78")
    for values in TABLE_ROWS:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0); paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs: _set_font(run, 9.5)
    _fixed_table(table, [2160, 3240, 3960])
    document.add_paragraph("Buenas prácticas recomendadas", style="Heading 1")
    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"; bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5); bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8); bullet.paragraph_format.line_spacing = 1.167
    for practice in PRACTICES: document.add_paragraph(practice, style="List Bullet")
    document.add_paragraph("Ejemplo práctico", style="Heading 1")
    document.add_paragraph(
        "Si una página tarda 8 segundos en cargar, primero mediría qué recursos ocupan más tiempo. Revisaría el peso de las imágenes, el trabajo de los scripts y las políticas de caché antes de asumir que existe una sola causa."
    )
    document.add_paragraph(
        "Después redimensionaría y comprimiría las imágenes. Mantendría en la carga inicial solo los scripts esenciales, cargaría los módulos secundarios bajo demanda y retiraría dependencias duplicadas o sin uso."
    )
    document.add_paragraph(
        "Por último configuraría caché para imágenes, estilos y scripts versionados. Volvería a medir en condiciones similares, comprobaría teclado y móvil, y repetiría los ajustes desde el cuello de botella más grande. Así buscaría una mejora visible sin romper funciones importantes."
    )
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("Planificación y Optimización Web"), 9, color="666666")
    document.core_properties.title = TITLE
    document.core_properties.author = "ARCHI"
    document.core_properties.subject = "Actividad académica de planificación y optimización web"
    document.save(output)


def validate(docx_path: Path, pdf_path: Path, json_path: Path) -> dict[str, object]:
    document = Document(docx_path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    joined = "\n".join(paragraphs)
    tables = document.tables
    pdf_valid = False; pdf_pages = 0; pdf_text = ""
    if pdf_path.is_file():
        reader = PdfReader(pdf_path); pdf_pages = len(reader.pages)
        pdf_text = "\n".join((page.extract_text() or "") for page in reader.pages)
        pdf_valid = pdf_pages > 0 and bool(pdf_text.strip()) and all((page.extract_text() or "").strip() for page in reader.pages)
    result = {
        "title_present": TITLE in joined,
        "introduction_present": "Introducción" in joined,
        "questions_answered": sum(label in joined for label, _, _ in QUESTIONS),
        "table_rows": max(0, len(tables[0].rows) - 1) if tables else 0,
        "best_practices_count": sum(practice in joined for practice in PRACTICES),
        "practical_example_present": "Ejemplo práctico" in joined,
        "generic_conclusion_absent": not any(value.casefold() == "conclusión" for value in paragraphs),
        "docx_valid": bool(paragraphs) and bool(tables),
        "pdf_valid": pdf_valid,
        "pdf_pages": pdf_pages,
        "pdf_text_characters": len(pdf_text),
    }
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--validate", action="store_true")
    args = parser.parse_args()
    docx = args.output_dir / "Actividad_Planificacion_Optimizacion_Web_ARCHI.docx"
    pdf = args.output_dir / "Actividad_Planificacion_Optimizacion_Web_ARCHI.pdf"
    report = args.output_dir / "ARCHI-ACTIVITY-VALIDATION.json"
    if not args.validate: build(docx)
    result = validate(docx, pdf, report)
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["docx_valid"] and (not args.validate or result["pdf_valid"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
