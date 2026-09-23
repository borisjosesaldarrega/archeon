from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from archeon.artifacts import ArtifactProvider, ArtifactSpec


class ArtifactVisualQualityV2Tests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.provider = ArtifactProvider()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_docx_uses_letter_page_and_plain_black_title(self) -> None:
        from docx import Document
        from docx.shared import Mm

        target = self.root / "report.docx"
        self.provider.create(ArtifactSpec(
            target, "docx", "Informe educativo", "Introducción clara.",
            [{"heading": "Contenido", "body": "Texto del informe."}],
            metadata={"style_profile": {"page_size": "Letter", "heading_color": "000000"}},
        ))
        document = Document(target)
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(str(document.styles["Title"].font.color.rgb), "000000")
        self.assertAlmostEqual(document.sections[0].page_width.mm, Mm(215.9).mm, places=1)

    def test_pdf_has_visual_structure_and_extractable_content(self) -> None:
        from pypdf import PdfReader

        target = self.root / "summary.pdf"
        result = self.provider.create(ArtifactSpec(
            target, "pdf", "Guía rápida", "Resumen visual para estudiantes.",
            [
                {"heading": "Aportes", "body": "Tutoría | práctica | accesibilidad"},
                {"heading": "Riesgos", "body": "Errores, sesgos y privacidad."},
            ],
        ))
        text = "\n".join(page.extract_text() or "" for page in PdfReader(target).pages)
        self.assertTrue(result.verified)
        self.assertIn(b"0.078 0.239 0.349 rg", target.read_bytes())
        self.assertIn("Guía rápida", text)
        self.assertIn("Aportes", text)

    def test_pptx_is_widescreen_and_not_an_office_default_layout(self) -> None:
        from pptx import Presentation

        target = self.root / "deck.pptx"
        self.provider.create(ArtifactSpec(
            target, "pptx", "IA en educación", slides=[
                {"title": "IA en educación", "body": "Aprender con criterio"},
                {"title": "Ventajas", "body": "Apoyo personalizado\nRetroalimentación rápida"},
            ],
        ))
        deck = Presentation(target)
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=2)
        self.assertEqual(deck.core_properties.comments, "ARCHI styled presentation v2")
        self.assertTrue(all(len(slide.shapes) >= 3 for slide in deck.slides))
        body = next(shape for shape in deck.slides[1].shapes if getattr(shape, "has_text_frame", False) and "Apoyo" in shape.text)
        self.assertTrue(all("<a:buChar" in paragraph._p.xml for paragraph in body.text_frame.paragraphs))

    def test_png_keeps_embedded_accessible_text(self) -> None:
        target = self.root / "infographic.png"
        result = self.provider.create(ArtifactSpec(
            target, "png", "Dos formas de aprender", metadata={"items": [
                {"title": "Practicar", "body": "Ejercicios adaptados."},
                {"title": "Revisar", "body": "Correcciones oportunas."},
            ]},
        ))
        self.assertTrue(result.verified)
        self.assertIn("Practicar", result.structure["content"])


if __name__ == "__main__":
    unittest.main()
